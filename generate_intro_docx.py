"""Generate Word version of project introduction."""
import re, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
style = doc.styles['Normal']
style.font.name = 'SimSun'
style.font.size = Pt(10.5)
style.font.color.rgb = RGBColor(0, 0, 0)
style.paragraph_format.space_after = Pt(3)
style.paragraph_format.line_spacing = 1.25

for lvl, sz in [(1, 16), (2, 13), (3, 11)]:
    s = doc.styles['Heading %d' % lvl]
    s.font.name = 'SimHei'
    s.font.size = Pt(sz)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)

with open('作品简介.md', 'r', encoding='utf-8') as f:
    lines = f.readlines()

for line in lines:
    line = line.rstrip()
    if not line.strip():
        doc.add_paragraph()
        continue
    if line.startswith('# ') and not line.startswith('## '):
        doc.add_heading(line[2:], level=1)
    elif line.startswith('## '):
        doc.add_heading(line[3:], level=2)
    elif line.startswith('!['):
        m = re.match(r'!\[(.*?)\]\((.*?)\)', line)
        if m:
            fp = os.path.join('report-assets', os.path.basename(m.group(2)))
            if os.path.exists(fp):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(fp, width=Inches(5.5))
    elif line.startswith('> '):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        r = p.add_run(line[2:])
        r.italic = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0, 0, 0)
    elif line.startswith('---'):
        continue
    else:
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', line)
        text = re.sub(r'`([^`]+)`', r'\1', text)
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.color.rgb = RGBColor(0, 0, 0)
        r.font.name = 'SimSun'

fp = '作品简介-v2.docx'
doc.save(fp)
print('Saved: %s (%d KB)' % (fp, os.path.getsize(fp) / 1024))
