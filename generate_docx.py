"""Convert design document markdown to formatted Word document."""
import re, os
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r'E:\zhuomian\c444\lz\vibe-sec-platform\系统设计文档-最终版.md'
DST = r'E:\zhuomian\c444\lz\vibe-sec-platform\系统设计文档-最终版-v2.docx'
IMG_DIR = r'E:\zhuomian\c444\lz\vibe-sec-platform\report-assets'

doc = Document()

# --- Style setup ---
style = doc.styles['Normal']
style.font.name = '宋体'
style.font.size = Pt(10.5)
style.font.color.rgb = RGBColor(0, 0, 0)
style.paragraph_format.space_after = Pt(4)
style.paragraph_format.line_spacing = 1.35
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

for level, (size, bold) in {1: (18, True), 2: (14, True), 3: (12, True)}.items():
    sname = f'Heading {level}'
    if sname in [s.name for s in doc.styles]:
        s = doc.styles[sname]
    else:
        s = doc.styles.add_style(sname, WD_STYLE_TYPE.PARAGRAPH)
    s.font.name = '黑体'
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    s.paragraph_format.space_before = Pt(12 if level > 1 else 20)
    s.paragraph_format.space_after = Pt(6)

with open(SRC, 'r', encoding='utf-8') as f:
    lines = f.readlines()

def add_paragraph(text, style='Normal', bold=False):
    p = doc.add_paragraph(style=style)
    if text.strip():
        run = p.add_run(text)
        run.bold = bold
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_code_block(code_lines):
    for line in code_lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        run = p.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0, 0, 0)

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers), style='Table Grid')
    # Header
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.name = '宋体'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        # Grey header bg
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'E8E8E8')
        cell._tc.get_or_add_tcPr().append(shading)
    # Data
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.rows[r+1].cells[c]
            cell.text = ''
            run = cell.paragraphs[0].add_run(str(val))
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.font.name = '宋体'
            run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def add_image(filename, caption=None):
    path = os.path.join(IMG_DIR, filename)
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(path, width=Inches(5.8))
        if caption:
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cap.add_run(caption)
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0, 0, 0)
            r.font.name = '宋体'
            r.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

# --- Parse markdown ---
i = 0
in_code = False
code_buf = []
in_table = False
table_buf = []

while i < len(lines):
    line = lines[i].rstrip()

    # Image
    img_match = re.match(r'!\[(.*?)\]\(report-assets/(.*?\.png)\)', line)
    if img_match:
        caption, filename = img_match.groups()
        add_image(filename, caption)
        i += 1
        continue

    # Code block
    if line.startswith('```'):
        if in_code:
            add_code_block(code_buf)
            code_buf = []
            in_code = False
        else:
            in_code = True
        i += 1
        continue
    if in_code:
        code_buf.append(line)
        i += 1
        continue

    # Table - collect lines until empty
    if line.startswith('|') and '|---' not in line:
        table_buf.append(line)
        i += 1
        # Check if next line starts table or is empty
        if i < len(lines) and not lines[i].strip().startswith('|'):
            # Parse accumulated table
            rows = []
            for tline in table_buf:
                cells = [c.strip() for c in tline.split('|')[1:-1]]
                rows.append(cells)
            if len(rows) > 1:
                # Skip separator row (---|---)
                data_rows = [r for r in rows[1:] if not all(re.match(r'^[-:]+$', c) for c in r)]
                if data_rows:
                    add_table(rows[0], data_rows)
            table_buf = []
        continue
    if table_buf and not line.startswith('|'):
        rows = []
        for tline in table_buf:
            cells = [c.strip() for c in tline.split('|')[1:-1]]
            rows.append(cells)
        if len(rows) > 1:
            data_rows = [r for r in rows[1:] if not all(re.match(r'^[-:]+$', c) for c in r)]
            if data_rows:
                add_table(rows[0], data_rows)
        table_buf = []
        # fall through to process current line

    # H1
    if line.startswith('# ') and not line.startswith('## '):
        add_paragraph(line[2:], 'Heading 1')
        i += 1
        continue

    # H2
    if line.startswith('## '):
        add_paragraph(line[3:], 'Heading 2')
        i += 1
        continue

    # Separator
    if line.strip() == '---':
        doc.add_paragraph()  # blank line
        i += 1
        continue

    # Blockquote
    if line.startswith('> '):
        text = line[2:]
        # Handle bold in blockquote
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.5)
        run = p.add_run(text)
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.italic = True
        run.font.name = '宋体'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        i += 1
        continue

    # Bold handling in normal text
    if line.strip():
        # Remove markdown bold markers for Word (Word uses actual bold formatting)
        text = line
        # Handle inline code
        text = re.sub(r'`([^`]+)`', r'\1', text)
        # Handle bold
        parts = re.split(r'(\*\*.*?\*\*)', text)
        if len(parts) > 1:
            p = doc.add_paragraph()
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.font.name = '宋体'
                    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                elif part.strip():
                    run = p.add_run(part)
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.font.name = '宋体'
                    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        else:
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.font.name = '宋体'
            run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    i += 1

# Save
doc.save(DST)
print(f'Saved: {DST}')
print(f'Size: {os.path.getsize(DST)/1024:.0f} KB')
